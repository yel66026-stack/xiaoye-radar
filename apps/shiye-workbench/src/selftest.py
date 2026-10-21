import csv
import json
from pathlib import Path
import sqlite3
import threading
import time
import traceback
from unittest.mock import patch

from engine import Engine, CARDS_SCRIPT, GATE_SCRIPT, browser_paths, likes_number, normalize_card
from store import Store, seed_demo


def wait_until(predicate, timeout=20):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return
        time.sleep(.05)
    raise AssertionError('Timed out waiting for expected state')


class FixtureEngine(Engine):
    """Real Playwright browser against intercepted, local-only fixtures."""
    def _ensure_browser(self):
        if self.context:
            return self.context.pages[0]
        from playwright.sync_api import sync_playwright
        self.pw = sync_playwright().start()
        self.context = self.pw.chromium.launch_persistent_context(str(self.home / 'fixture_profile'),
            executable_path=next(iter(browser_paths().values())), headless=True)
        self.context.set_default_navigation_timeout(5000)
        cards = []
        for identity, token, title in [('a', 'token1', '中文标题，含逗号'), ('a', 'changed', '重复链接'), ('b', 'token2', '=测试公式防护')]:
            cards.append(f'<section class="note-item"><a class="cover" href="https://www.xiaohongshu.com/explore/{identity}?xsec_token={token}">封面</a><div class="title">{title}</div><div class="author-wrapper"><div class="name">测试作者</div></div><div class="like-wrapper"><span class="count">1.2万</span></div></section>')
        self.fixture_html = '<meta charset="utf-8">' + ''.join(cards)
        def route_handler(route):
            html = self.fixture_html if '/search_result?' in route.request.url else '<meta charset="utf-8"><div id="detail-desc">测试正文\n第二行内容 #标签</div>'
            route.fulfill(status=200, content_type='text/html; charset=utf-8', body=html)
        self.context.route('**/*', route_handler)
        self.browser_status = 'open'
        return self.context.pages[0] if self.context.pages else self.context.new_page()

    def _checkpoint(self, seconds=0):
        return super()._checkpoint(min(seconds, .02))


def run_tests(directory):
    directory.mkdir(parents=True, exist_ok=True)
    report = {'passed': False, 'checks': [], 'scope': 'Offline fixtures; no real Xiaohongshu login or live data collection.'}
    engine = None
    api = None
    try:
        store = Store(directory / 'workspace.sqlite3', directory / 'exports')
        seed_demo(store)
        assert store.stats()['total'] == 12
        assert store.stats()['unique_notes'] == 12
        assert store.notes({'search': '城市周末'})['total'] == 6
        assert store.notes({'min_likes': 5000})['total'] == 2
        assert store.notes({'search': '%_\\'})['total'] == 0
        assert len(store.notes({'ids': [1,2]})['rows']) == 2
        assert store.notes({'sort': 'likes'})['rows'][0]['likes'] == 7260
        assert likes_number('1.2万') == 12000 and likes_number('3.4k') == 3400 and likes_number('点赞') == 0
        assert normalize_card({'link': 'javascript:alert(1)'}, '关键词') is None
        assert normalize_card({'link': 'https://evil.example/explore/a'}, '关键词') is None
        report['checks'].append('SQLite transactions, search escaping, filters, sorting, selected rows, link validation and Chinese like counts')

        backup = store.backup(directory / 'backups')
        with sqlite3.connect(backup) as db:
            assert db.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
            assert db.execute('SELECT COUNT(*) FROM notes').fetchone()[0] == 12
        pending = store.create_task('恢复测试', 10, 'fast')
        store.update_task(pending, status='running')
        reopened = Store(store.path, directory / 'exports')
        assert reopened.task(pending)['status'] == 'interrupted'
        report['checks'].append('Online database backup and abnormal-exit task recovery')

        # Deep mode through the production task worker, with a real browser.
        engine = FixtureEngine(store, directory)
        identity = store.create_task('测试/旅行', 2, 'deep')
        engine.submit(identity)
        wait_until(lambda: store.task(identity)['status'] in ('completed','error'), 25)
        task = store.task(identity)
        assert task['status'] == 'completed', task
        notes = store.notes({'task_id': identity})['rows']
        assert len(notes) == 2 and all(n['detail_status']=='full' for n in notes)
        assert all('第二行内容' in n['description'] for n in notes)
        assert any('xsec_token=token1' in n['link'] for n in notes)
        report['checks'].append('Playwright engine: deep-mode collection, ID deduplication, signed links, body extraction and completion')

        selected = [n['id'] for n in notes]
        for fmt in ('csv','xlsx','json'):
            result = store.export(directory / 'exports', fmt, {'ids': selected}, prefix='测试/非法:*文件名')
            assert Path(result['path']).exists() and result['count'] == 2
            if fmt == 'csv':
                assert Path(result['path']).read_bytes().startswith(b'\xef\xbb\xbf')
                with open(result['path'], encoding='utf-8-sig', newline='') as file:
                    rows = list(csv.reader(file))
                assert len(rows)==3 and any(row[1].startswith("'=") for row in rows[1:])
            elif fmt == 'xlsx':
                from openpyxl import load_workbook
                book = load_workbook(result['path'])
                sheet = book.active
                assert sheet.max_row == 3 and sheet.freeze_panes == 'A2'
                assert all(sheet.cell(row,2).data_type != 'f' for row in (2,3))
                book.close()
            else:
                assert len(json.loads(Path(result['path']).read_text(encoding='utf-8'))) == 2
        report['checks'].append('Excel/CSV/JSON exports: row scope, Chinese text, formula defense, workbook style and atomic files')

        # Production checkpoint/pause/continue/stop states, no browser navigation required.
        wait_until(lambda: engine.active is None)
        gate_started = threading.Event()
        def controlled_collect(identity):
            gate_started.set()
            while True:
                engine._checkpoint(.05)
        with patch.object(engine, '_collect', side_effect=controlled_collect):
            paused = store.create_task('暂停与停止', 20, 'fast')
            engine.submit(paused)
            assert gate_started.wait(3)
            engine.action(paused,'pause')
            assert store.task(paused)['status'] == 'paused'
            engine.action(paused,'resume')
            assert store.task(paused)['status'] == 'running'
            engine.action(paused,'stop')
            wait_until(lambda: store.task(paused)['status']=='stopped')
            wait_until(lambda: engine.active is None)
        report['checks'].append('Task pause, resume, responsive cancellation and retained database records')
        engine.shutdown()
        engine = None

        from desktop import DesktopApi
        api = DesktopApi(directory / 'api')
        snap = api.snapshot()
        assert snap['ok'] and snap['data']['stats']['total'] == 0
        assert api.create_tasks({'keywords':['test'],'target':0,'mode':'fast'})['ok'] is False
        assert api.save_settings({'workspace':'','export_dir':'relative'})['ok'] is False
        assert api.toggle_demo()['ok']
        assert api.snapshot()['data']['stats']['total'] == 12
        assert api.create_tasks({'keywords':['test'],'target':10,'mode':'fast'})['ok'] is False
        assert api.open_note(1)['ok'] is False
        assert api.toggle_demo()['ok']
        assert api.snapshot()['data']['stats']['total'] == 0
        report['checks'].append('Desktop bridge: validation, demo separation, blocked demo collection and isolated local workspaces')
        from condition_test import run_condition_checks
        report['checks'].extend(run_condition_checks(directory/'conditions'))
        report['passed'] = True
    except Exception:
        report['error'] = traceback.format_exc()
    finally:
        if engine:
            engine.shutdown()
        if api:
            api._engine.shutdown()
        (directory / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if report['passed'] else 1
