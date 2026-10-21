import csv
import json
import re
import sqlite3
import threading
import uuid
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from rules import split_terms, normalize_region, SHANGHAI, fold_text

VERSION = '2.2.2'
ACTIVE = ('running', 'paused', 'waiting')
TERMINAL = ('completed', 'partial', 'stopped', 'error', 'interrupted')
NOTE_COLUMNS = ['keyword','title','description','author','likes_text','likes','ip_location','ip_source','published_at',
                'published_raw','published_precision','published_source','time_kind','note_type','matched_terms','link','collected_at','detail_status']
HEADERS = ['关键词','笔记标题','正文 / 简介','作者','点赞原文','点赞数值','IP 属地','属地来源','发布时间',
           '页面时间原文','时间精度','发布时间来源','时间类型','内容类型','命中词汇','笔记链接','采集时间','正文状态']
EXTRA_NOTE_DEFAULTS = {'ip_location':'','ip_source':'','ip_raw':'','published_at':'','published_raw':'',
    'published_precision':'unknown','published_source':'','time_kind':'unknown','note_type':'unknown',
    'author_link':'','matched_terms':'[]','likes_known':1}


def now():
    return datetime.now().astimezone().isoformat(timespec='seconds')


def safe_name(text):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(text)).strip(' .')[:60] or '采集结果'


def safe_cell(value):
    if isinstance(value, (int, float)):
        return value
    text = str(value or '')
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r')) else text


class Store:
    def __init__(self, path, export_dir):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.migration_backup = None
        if self.path.exists() and self.path.stat().st_size:
            with self.connection() as db:
                columns = {row[1] for row in db.execute('PRAGMA table_info(notes)')}
                if columns and 'ip_location' not in columns:
                    backup = self.path.with_name(self.path.stem + f'.pre-v2.1-{datetime.now():%Y%m%d_%H%M%S_%f}.sqlite3')
                    target = sqlite3.connect(backup)
                    try:
                        db.backup(target)
                    finally:
                        target.close()
                    self.migration_backup = str(backup)
        with self.connection() as db:
            db.executescript('''
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS tasks(
              id TEXT PRIMARY KEY, keyword TEXT NOT NULL, target INTEGER NOT NULL,
              mode TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'queued',
              created_at TEXT NOT NULL, started_at TEXT, ended_at TEXT,
              elapsed REAL NOT NULL DEFAULT 0, message TEXT NOT NULL DEFAULT '',
              scroll_y INTEGER NOT NULL DEFAULT 0);
            CREATE TABLE IF NOT EXISTS notes(
              id INTEGER PRIMARY KEY, task_id TEXT NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
              note_id TEXT NOT NULL, keyword TEXT NOT NULL, title TEXT NOT NULL,
              description TEXT NOT NULL, author TEXT NOT NULL, likes_text TEXT NOT NULL,
              likes INTEGER NOT NULL DEFAULT 0, link TEXT NOT NULL,
              collected_at TEXT NOT NULL, detail_status TEXT NOT NULL DEFAULT 'title',
              UNIQUE(task_id,note_id));
            CREATE INDEX IF NOT EXISTS ix_notes_task ON notes(task_id);
            CREATE INDEX IF NOT EXISTS ix_notes_time ON notes(collected_at);
            CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY,task_id TEXT,level TEXT NOT NULL,message TEXT NOT NULL,created_at TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS word_libraries(id TEXT PRIMARY KEY,name TEXT NOT NULL,description TEXT NOT NULL DEFAULT '',
                search_terms TEXT NOT NULL,include_words TEXT NOT NULL,exclude_words TEXT NOT NULL,include_mode TEXT NOT NULL,
                created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS candidates(task_id TEXT NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,note_id TEXT NOT NULL,
                outcome TEXT NOT NULL,reason TEXT NOT NULL,scanned_at TEXT NOT NULL,attempts INTEGER NOT NULL DEFAULT 1,
                PRIMARY KEY(task_id,note_id));
            ''')
            task_columns = {row[1] for row in db.execute('PRAGMA table_info(tasks)')}
            for name in ('criteria_json','library_snapshot'):
                if name not in task_columns:
                    db.execute(f"ALTER TABLE tasks ADD COLUMN {name} TEXT NOT NULL DEFAULT '{{}}'")
            note_columns = {row[1] for row in db.execute('PRAGMA table_info(notes)')}
            for key, default in EXTRA_NOTE_DEFAULTS.items():
                if key not in note_columns:
                    definition = f'INTEGER NOT NULL DEFAULT {default}' if isinstance(default,int) else "TEXT NOT NULL DEFAULT '" + default.replace("'","''") + "'"
                    db.execute(f'ALTER TABLE notes ADD COLUMN {key} {definition}')
            db.execute('CREATE INDEX IF NOT EXISTS ix_notes_region ON notes(ip_location)')
            db.execute('CREATE INDEX IF NOT EXISTS ix_notes_published ON notes(published_at)')
            db.execute("INSERT OR REPLACE INTO meta VALUES('schema','2')")
            db.execute("INSERT OR IGNORE INTO settings VALUES('export_dir',?)", (json.dumps(str(export_dir), ensure_ascii=False),))
            db.execute("UPDATE tasks SET status='interrupted',message='上次程序未正常退出，已保留数据，可继续任务。',ended_at=? WHERE status IN ('running','paused','waiting','queued')", (now(),))

    @contextmanager
    def connection(self):
        with self.lock:
            db = sqlite3.connect(str(self.path), timeout=15)
            db.row_factory = sqlite3.Row
            db.create_function('fold_text',1,fold_text,deterministic=True)
            db.execute('PRAGMA foreign_keys=ON')
            try:
                yield db
                db.commit()
            except Exception:
                db.rollback()
                raise
            finally:
                db.close()

    def settings(self):
        value = {'workspace': '我的内容工作区', 'browser': 'auto', 'interval': 4, 'detail_interval': 2.5,
                 'auto_export': True, 'export_format': 'xlsx', 'export_dir': str(self.path.parent.parent / '采集结果')}
        with self.connection() as db:
            value.update({row['key']: json.loads(row['value']) for row in db.execute('SELECT * FROM settings')})
        return value

    def save_settings(self, values):
        with self.connection() as db:
            db.executemany('INSERT OR REPLACE INTO settings VALUES(?,?)',
                           [(key, json.dumps(value, ensure_ascii=False)) for key, value in values.items()])

    def create_task(self, keyword, target, mode, criteria=None, library_snapshot=None):
        identity = uuid.uuid4().hex
        with self.connection() as db:
            db.execute('INSERT INTO tasks(id,keyword,target,mode,created_at,criteria_json,library_snapshot) VALUES(?,?,?,?,?,?,?)',
                       (identity, keyword, target, mode, now(),json.dumps(criteria or {},ensure_ascii=False),json.dumps(library_snapshot or {},ensure_ascii=False)))
        self.event(identity, 'info', f'任务已创建：{keyword}，目标 {target} 条')
        return identity

    def task(self, identity):
        with self.connection() as db:
            row = db.execute('SELECT t.*,(SELECT COUNT(*) FROM notes n WHERE n.task_id=t.id) AS count FROM tasks t WHERE id=?', (identity,)).fetchone()
            if not row:
                raise ValueError('任务不存在')
            return self._task_dict(row, db)

    def tasks(self):
        with self.connection() as db:
            return [self._task_dict(row,db) for row in db.execute('SELECT t.*,(SELECT COUNT(*) FROM notes n WHERE n.task_id=t.id) AS count FROM tasks t ORDER BY created_at DESC,rowid DESC LIMIT 500')]

    @staticmethod
    def _task_dict(row, db):
        value = dict(row)
        value['criteria'] = json.loads(value.pop('criteria_json') or '{}')
        value['library'] = json.loads(value.pop('library_snapshot') or '{}')
        counts = list(db.execute('SELECT outcome,reason,COUNT(*) amount FROM candidates WHERE task_id=? GROUP BY outcome,reason',(value['id'],)))
        value['skipped'] = sum(r['amount'] for r in counts if r['outcome']!='matched')
        value['scanned'] = value['count'] + value['skipped']
        value['skip_reasons'] = {r['reason']:r['amount'] for r in counts if r['outcome']!='matched'}
        return value

    def update_task(self, identity, **values):
        allowed = {'status', 'started_at', 'ended_at', 'elapsed', 'message', 'scroll_y'}
        if not values or not set(values) <= allowed:
            raise ValueError('无效任务字段')
        with self.connection() as db:
            db.execute(f"UPDATE tasks SET {','.join(key+'=?' for key in values)} WHERE id=?", (*values.values(), identity))

    def event(self, identity, level, message):
        with self.connection() as db:
            db.execute('INSERT INTO events(task_id,level,message,created_at) VALUES(?,?,?,?)', (identity, level, str(message)[:2000], now()))

    def recent_events(self, identity=None):
        with self.connection() as db:
            query = 'SELECT * FROM events'
            args = ()
            if identity:
                query += ' WHERE task_id=?'
                args = (identity,)
            return [dict(row) for row in db.execute(query + ' ORDER BY id DESC LIMIT 80', args)]

    def add_note(self, identity, note):
        with self.connection() as db:
            return self._insert_note(db,identity,note)

    @staticmethod
    def _insert_note(db, identity, note):
        row = {key:note.get(key,default) for key,default in EXTRA_NOTE_DEFAULTS.items()}
        row.update({key:note.get(key,'') for key in ['note_id','keyword','title','description','author','likes_text','link']})
        row.update(task_id=identity,likes=note.get('likes',0),collected_at=note.get('collected_at',now()),detail_status=note.get('detail_status','title'))
        if isinstance(row['matched_terms'],list):
            row['matched_terms']=json.dumps(row['matched_terms'],ensure_ascii=False)
        before=db.total_changes
        db.execute('INSERT OR IGNORE INTO notes('+','.join(row)+') VALUES('+','.join('?' for _ in row)+')',list(row.values()))
        return db.total_changes>before

    def record_candidate(self, identity, note_id, outcome, reason='', note=None):
        if outcome not in ('matched','filtered','unknown'):
            raise ValueError('无效扫描结果')
        with self.connection() as db:
            added = self._insert_note(db,identity,note) if note is not None else False
            db.execute('''INSERT INTO candidates(task_id,note_id,outcome,reason,scanned_at) VALUES(?,?,?,?,?)
                ON CONFLICT(task_id,note_id) DO UPDATE SET outcome=excluded.outcome,reason=excluded.reason,
                scanned_at=excluded.scanned_at,attempts=candidates.attempts+1''',(identity,note_id,outcome,reason,now()))
            return added

    def seen_ids(self, identity):
        with self.connection() as db:
            return {row[0] for row in db.execute("SELECT note_id FROM notes WHERE task_id=? UNION SELECT note_id FROM candidates WHERE task_id=? AND outcome!='unknown'", (identity,identity))}

    def libraries(self):
        with self.connection() as db:
            return [self._library_dict(row) for row in db.execute('SELECT * FROM word_libraries ORDER BY updated_at DESC,rowid DESC')]

    @staticmethod
    def _library_dict(row):
        value=dict(row)
        for key in ('search_terms','include_words','exclude_words'):
            value[key]=json.loads(value[key])
        return value

    def library(self, identity):
        with self.connection() as db:
            row=db.execute('SELECT * FROM word_libraries WHERE id=?',(identity,)).fetchone()
            if not row:
                raise ValueError('词汇库不存在或已删除。')
            return self._library_dict(row)

    def save_library(self, payload):
        name=str(payload.get('name') or '').strip()
        if not name or len(name)>40:
            raise ValueError('词库名称需要 1–40 个字符。')
        values={key:split_terms(payload.get(key),30 if key=='search_terms' else 300,80 if key=='search_terms' else 100)
                for key in ('search_terms','include_words','exclude_words')}
        if not any(values.values()):
            raise ValueError('词库至少需要一条搜索词、包含词或排除词。')
        mode=payload.get('include_mode','any')
        if mode not in ('any','all'):
            raise ValueError('请选择有效的包含词匹配方式。')
        identity=payload.get('id') or uuid.uuid4().hex
        with self.connection() as db:
            if payload.get('id') and not db.execute('SELECT id FROM word_libraries WHERE id=?',(identity,)).fetchone():
                raise ValueError('待编辑词汇库不存在。')
            if not payload.get('id') and db.execute('SELECT COUNT(*) FROM word_libraries').fetchone()[0]>=100:
                raise ValueError('最多保存 100 个词汇库。')
            db.execute('''INSERT INTO word_libraries(id,name,description,search_terms,include_words,exclude_words,include_mode,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name,description=excluded.description,
                search_terms=excluded.search_terms,include_words=excluded.include_words,exclude_words=excluded.exclude_words,
                include_mode=excluded.include_mode,updated_at=excluded.updated_at''',
                (identity,name,str(payload.get('description') or '')[:300],*(json.dumps(values[key],ensure_ascii=False) for key in ('search_terms','include_words','exclude_words')),mode,now(),now()))
        return self.library(identity)

    def delete_library(self, identity):
        with self.connection() as db:
            db.execute('DELETE FROM word_libraries WHERE id=?',(identity,))

    def notes(self, filters=None, all_rows=False):
        filters = filters or {}
        clauses, args = [], []
        if filters.get('task_id'):
            clauses.append('task_id=?')
            args.append(str(filters['task_id']))
        if filters.get('ids'):
            identities = [int(value) for value in filters['ids']][:1000]
            clauses.append('id IN (' + ','.join('?' for _ in identities) + ')')
            args.extend(identities)
        if filters.get('search'):
            clauses.append('(title LIKE ? ESCAPE \'\\\' OR author LIKE ? ESCAPE \'\\\' OR keyword LIKE ? ESCAPE \'\\\')')
            search = str(filters['search'])[:100].replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
            args.extend(['%' + search + '%'] * 3)
        if filters.get('min_likes') is not None and filters.get('min_likes')!='':
            clauses.append('likes>=?')
            args.append(max(0, int(filters['min_likes'])))
            clauses.append('likes_known=1')
        if filters.get('max_likes') is not None and filters.get('max_likes')!='':
            clauses.extend(['likes<=?','likes_known=1'])
            args.append(max(0,int(filters['max_likes'])))
        regions = [normalize_region(value) for value in split_terms(filters.get('regions'),50,40)]
        if any(not region for region in regions):
            raise ValueError('请选择明确的 IP 属地地区。')
        if regions:
            clauses.append('ip_location IN ('+','.join('?' for _ in regions)+')')
            args.extend(regions)
        if regions and filters.get('ip_source') in ('note','author'):
            clauses.append('ip_source=?')
            args.append(filters['ip_source'])
        for field,comparison in (('date_from','>='),('date_to','<=')):
            if filters.get(field):
                day=datetime.strptime(str(filters[field]),'%Y-%m-%d').date().isoformat()
                clauses.extend([f'published_at{comparison}?',"time_kind='published'"])
                args.append(day)
        if filters.get('note_type') in ('image','video'):
            clauses.append('note_type=?')
            args.append(filters['note_type'])
        for key,negate in (('authors',False),('exclude_authors',True)):
            words=split_terms(filters.get(key),100,80)
            if words:
                clauses.append("fold_text(author) NOT IN ('','未知作者')")
                clauses.append('fold_text(author) '+('NOT IN' if negate else 'IN')+' ('+','.join('fold_text(?)' for _ in words)+')')
                args.extend(words)
        include=split_terms(filters.get('include_words'))
        exclude=split_terms(filters.get('exclude_words'))
        scope=filters.get('text_scope','both')
        if scope not in ('title','body','both'):
            raise ValueError('无效匹配范围')
        expression={'title':'title','body':'description','both':"title||char(10)||description"}[scope]
        if (include or exclude) and scope!='title':
            clauses.append("detail_status='full'")
        if include:
            joiner=' AND ' if filters.get('include_mode')=='all' else ' OR '
            clauses.append('('+joiner.join(f'instr(fold_text({expression}),fold_text(?))>0' for _ in include)+')')
            args.extend(include)
        if exclude:
            clauses.extend(f'instr(fold_text({expression}),fold_text(?))=0' for _ in exclude)
            args.extend(exclude)
        if filters.get('detail') == 'full':
            clauses.append("detail_status='full'")
        where = (' WHERE ' + ' AND '.join(clauses)) if clauses else ''
        order = {'likes': 'likes DESC,id DESC', 'oldest': 'id ASC'}.get(filters.get('sort'), 'id DESC')
        with self.connection() as db:
            total = db.execute('SELECT COUNT(*) FROM notes' + where, args).fetchone()[0]
            query = 'SELECT * FROM notes' + where + ' ORDER BY ' + order
            if not all_rows:
                limit = 30
                page = max(1, int(filters.get('page', 1)))
                query += ' LIMIT ? OFFSET ?'
                args.extend([limit, (page - 1) * limit])
            rows = [dict(row) for row in db.execute(query, args)]
            return {'total': total, 'rows': rows}

    def note(self, identity):
        with self.connection() as db:
            row = db.execute('SELECT * FROM notes WHERE id=?', (int(identity),)).fetchone()
            if not row:
                raise ValueError('笔记不存在')
            return dict(row)

    def stats(self):
        with self.connection() as db:
            row = db.execute('''SELECT COUNT(*) total,COUNT(DISTINCT note_id) unique_notes,
                COALESCE(SUM(likes),0) likes,COALESCE(SUM(detail_status='full'),0) full_notes FROM notes''').fetchone()
            result = dict(row)
            result['tasks'] = db.execute('SELECT COUNT(*) FROM tasks').fetchone()[0]
            today = datetime.now().astimezone().date().isoformat()
            result['today'] = db.execute('SELECT COUNT(*) FROM notes WHERE collected_at>=?', (today,)).fetchone()[0]
            result['keywords'] = [dict(row) for row in db.execute('SELECT keyword,COUNT(*) count,COALESCE(SUM(likes),0) likes FROM notes GROUP BY keyword ORDER BY count DESC LIMIT 6')]
            return result

    def export(self, folder, fmt, filters=None, prefix='拾页'):
        if fmt not in ('csv', 'xlsx', 'json'):
            raise ValueError('不支持的导出格式')
        rows = self.notes(filters, all_rows=True)['rows']
        if not rows:
            raise ValueError('没有符合条件的记录可导出')
        folder = Path(folder)
        folder.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
        path = folder / f'{safe_name(prefix)}_{stamp}.{fmt}'
        temporary = path.with_suffix(path.suffix + '.tmp')
        try:
            if fmt == 'json':
                exported=[{key:json.loads(row[key]) if key=='matched_terms' else row[key] for key in NOTE_COLUMNS} for row in rows]
                temporary.write_text(json.dumps(exported, ensure_ascii=False, indent=2), encoding='utf-8')
            elif fmt == 'csv':
                with temporary.open('x', encoding='utf-8-sig', newline='') as file:
                    writer = csv.writer(file)
                    writer.writerow(HEADERS)
                    writer.writerows([[safe_cell(self._export_cell(row,key)) for key in NOTE_COLUMNS] for row in rows])
            else:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
                from openpyxl.worksheet.table import Table, TableStyleInfo
                book = Workbook()
                sheet = book.active
                sheet.title = '笔记数据'
                sheet.append(HEADERS)
                for row in rows:
                    sheet.append([safe_cell(self._export_cell(row,key)) for key in NOTE_COLUMNS])
                sheet.freeze_panes = 'A2'
                for cell in sheet[1]:
                    cell.font = Font(name='微软雅黑', color='FFFFFF', bold=True)
                    cell.fill = PatternFill('solid', fgColor='1D2839')
                    cell.alignment = Alignment(vertical='center')
                sheet.row_dimensions[1].height = 30
                widths = [20,48,68,22,14,14,16,18,18,28,20,24,18,16,32,60,28,18]
                from openpyxl.utils import get_column_letter
                for index, width in enumerate(widths, 1):
                    sheet.column_dimensions[get_column_letter(index)].width = width
                table = Table(displayName='Notes', ref=f'A1:{get_column_letter(len(NOTE_COLUMNS))}{len(rows)+1}')
                table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
                sheet.add_table(table)
                book.save(temporary)
                book.close()
            temporary.replace(path)
        finally:
            if temporary.exists():
                temporary.unlink()
        return {'path': str(path), 'count': len(rows), 'format': fmt}

    @staticmethod
    def _export_cell(row,key):
        if key=='matched_terms':
            return '、'.join(json.loads(row[key]))
        mapping={'ip_source':{'note':'笔记公开属地','author':'作者主页公开属地'},
                 'published_source':{'page':'页面日期字段','page-meta':'页面发布时间元数据','json-ld':'页面结构化发布时间','demo':'演示数据'},
                 'note_type':{'image':'图文','video':'视频','unknown':'未读取'},
                 'time_kind':{'published':'发布时间','updated':'编辑时间（非发布时间）','unknown':'未读取'},
                 'published_precision':{'day':'页面完整日期','timestamp':'页面时间戳','relative':'相对时间换算','inferred_year':'按页面日期推算年份','unknown':'未读取'}}
        return mapping.get(key,{}).get(row[key],row[key])

    def backup(self, folder):
        folder = Path(folder)
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f'拾页_工作区备份_{datetime.now():%Y%m%d_%H%M%S_%f}.sqlite3'
        with self.connection() as source:
            dest = sqlite3.connect(path)
            try:
                source.backup(dest)
            finally:
                dest.close()
        return str(path)


def seed_demo(store):
    if store.tasks():
        _upgrade_demo(store)
        return
    samples = [
      ('城市周末', '把周末还给自己：6 条城市漫游路线', '慢游编辑部', 2860),
      ('城市周末', '不用请假，上海周边两天一夜这样玩', '阿柚的周末', 1842),
      ('城市周末', '雨天也想出门｜这几家小店值得收藏', '一颗橘子', 936),
      ('城市周末', '沿着河边走一走，找到城市的另一面', '走路观察员', 725),
      ('城市周末', '人均 200 的周末松弛感计划', '小满生活志', 3210),
      ('城市周末', '给工作按下暂停：我的周末半日清单', '椰子不加糖', 568),
      ('咖啡探店', '藏在老街里的咖啡店，我愿意再去十次', '拿铁研究所', 5620),
      ('咖啡探店', '这杯手冲，让阴天变得明亮了', '豆子和日常', 1580),
      ('咖啡探店', '周末咖啡地图｜安静、有光、好好喝', '米可的地图', 2100),
      ('咖啡探店', '不是网红店，是会想念的街角咖啡馆', '阿森的口袋', 482),
      ('家居灵感', '12 平米也能拥有的阅读角', '住进好天气', 7260),
      ('家居灵感', '小改动，大幸福：我的书桌整理记录', '木木的房间', 3510),
    ]
    for keyword in ('城市周末', '咖啡探店', '家居灵感'):
        matching = [row for row in samples if row[0] == keyword]
        identity = store.create_task(keyword, len(matching), 'deep')
        for index, (_, title, author, likes) in enumerate(matching):
            store.add_note(identity, {'note_id': f'demo-{keyword}-{index}', 'keyword': keyword, 'title': title,
             'description': f'【演示数据 · 非真实采集】\n\n{title}\n\n这是一条用于体验结果预览、筛选和导出的示例笔记。\n\n生活的灵感，常常藏在不经意的日常里。留一点时间，去观察、去记录，整理属于自己的内容清单。\n\n#生活方式 #灵感记录 #周末计划',
             'author': author, 'likes_text': f'{likes:,}', 'likes': likes, 'link': '', 'detail_status': 'full'})
        store.update_task(identity, status='completed', ended_at=now(), elapsed=48 + len(matching)*8, message='演示任务，数据仅供体验')
    store.event(None, 'info', '演示数据与真实工作区完全隔离，可以体验筛选、详情和导出。')
    _upgrade_demo(store)


def _upgrade_demo(store):
    with store.connection() as db:
        rows=list(db.execute("SELECT id FROM notes WHERE note_id LIKE 'demo-%' ORDER BY id"))
        today=datetime.now(SHANGHAI).date()
        for index,row in enumerate(rows):
            region=('北京','上海','广东')[index%3]
            db.execute("""UPDATE notes SET ip_location=?,ip_source='note',ip_raw=?,published_at=?,
               published_raw=?,published_precision='day',published_source='demo',time_kind='published',note_type=?
               WHERE id=? AND ip_location='' AND published_at=''""",
               (region,'IP属地：'+region,(today-timedelta(days=index)).isoformat(),(today-timedelta(days=index)).isoformat(),
                'video' if index%3==0 else 'image',row['id']))
    if not store.libraries():
        store.save_library({'name':'城市生活监测 · 演示','description':'演示词库，可编辑和导出，不会影响真实工作区。',
            'search_terms':['城市周末','咖啡探店'],'include_words':['周末','咖啡'],'exclude_words':['广告','招募'],'include_mode':'any'})
